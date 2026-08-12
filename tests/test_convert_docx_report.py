from pathlib import Path

from docx import Document
from PIL import Image

from tools.convert_docx_report import convert_docx


def test_convert_docx_preserves_headings_tables_and_images(tmp_path: Path) -> None:
    image_path = tmp_path / "source.png"
    Image.new("RGB", (40, 20), "red").save(image_path)

    source = tmp_path / "report.docx"
    document = Document()
    document.add_heading("Analiza próbna", level=1)
    document.add_paragraph("Treść z wartością 65 dB(A).")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Metryka"
    table.cell(0, 1).text = "Wartość"
    table.cell(1, 0).text = "Populacja"
    table.cell(1, 1).text = "100"
    document.add_picture(str(image_path))
    document.add_paragraph("Rysunek 1. Mapa próbna")
    document.save(source)

    output = tmp_path / "report.md"
    assets = tmp_path / "assets"
    result = convert_docx(source, output, assets, image_prefix="figure")

    markdown = output.read_text(encoding="utf-8")
    assert "# Analiza próbna" in markdown
    assert "| Metryka | Wartość |" in markdown
    assert "| Populacja | 100 |" in markdown
    assert "![Rysunek 1. Mapa próbna](assets/figure-001.png)" in markdown
    assert str(tmp_path) not in markdown
    assert result.table_count == 1
    assert result.image_count == 1


def test_convert_docx_can_use_semantic_image_names(tmp_path: Path) -> None:
    image_path = tmp_path / "source.png"
    Image.new("RGB", (40, 20), "blue").save(image_path)
    source = tmp_path / "report.docx"
    document = Document()
    document.add_picture(str(image_path))
    document.add_paragraph("")
    document.add_paragraph("Mapa 7. Podwójne ryzyko powodziowe")
    document.save(source)

    output = tmp_path / "report.md"
    convert_docx(source, output, tmp_path / "assets", semantic_names=True)

    markdown = output.read_text(encoding="utf-8")
    assert "assets/mapa-7-podwojne-ryzyko-powodziowe.png" in markdown


def test_convert_docx_removes_manual_toc_and_heading_formatting(tmp_path: Path) -> None:
    source = tmp_path / "report.docx"
    document = Document()
    title = document.add_paragraph()
    title.add_run("ANALIZA KIERUNKOWA").bold = True
    toc = document.add_heading("Spis treści", level=1)
    toc.runs[0].bold = True
    document.add_paragraph("Streszczenie 3")
    heading = document.add_heading("Streszczenie", level=1)
    heading.runs[0].bold = True
    document.add_paragraph("Treść raportu")
    document.save(source)

    output = tmp_path / "report.md"
    convert_docx(source, output, tmp_path / "assets")

    markdown = output.read_text(encoding="utf-8")
    assert markdown.startswith("# ANALIZA KIERUNKOWA\n")
    assert "Spis treści" not in markdown
    assert "Streszczenie 3" not in markdown
    assert "# Streszczenie\n" in markdown
    assert "# **Streszczenie**" not in markdown


def test_publication_mode_adds_version_and_machine_readable_links(tmp_path: Path) -> None:
    source = tmp_path / "report.docx"
    document = Document()
    document.add_paragraph("ANALIZA KIERUNKOWA")
    document.add_heading("Załączniki", level=1)
    document.save(source)

    output = tmp_path / "publication" / "report.md"
    convert_docx(
        source,
        output,
        tmp_path / "publication" / "assets",
        publication_metadata=True,
    )

    markdown = output.read_text(encoding="utf-8")
    assert "Wersja danych: `2026.05`" in markdown
    assert "Rozszerzenie CAP/IoT wyeksportowano 6 lipca 2026 r." in markdown
    assert "[A04 - gminy](tables/A04_gminy.csv)" in markdown
    assert "[manifest danych](../data/manifests/data-manifest.json)" in markdown
